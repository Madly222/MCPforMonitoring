"""
Posta Moldovei Invoice Generator
Connects to billing DB, queries clients, generates XLSX, uploads to FTP.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional
import paramiko
from loguru import logger

try:
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent.parent.parent / ".env")
except ImportError:
    pass


BILLING_CONFIG = {
    "ssh_host": os.getenv("INVOICE_SSH_HOST", "172.22.22.60"),
    "ssh_user": os.getenv("INVOICE_SSH_USER", "root"),
    "ssh_password": os.getenv("INVOICE_SSH_PASSWORD", ""),
    "ssh_port": int(os.getenv("INVOICE_SSH_PORT", 22)),
    "db_name": os.getenv("BILLING_DB_NAME", "billing"),
    "db_user": os.getenv("BILLING_DB_USER", "billing"),
    "db_password": os.getenv("BILLING_DB_PASSWORD", ""),
}

FTP_CONFIG = {
    "host": os.getenv("INVOICE_FTP_HOST", "ftp2.posta.md"),
    "user": os.getenv("INVOICE_FTP_USER", ""),
    "password": os.getenv("INVOICE_FTP_PASSWORD", ""),
    "path": os.getenv("INVOICE_FTP_PATH", "/In/"),
}

CONFIG_DIR = Path(__file__).parent.parent.parent / "config"
CLIENTS_FILE = CONFIG_DIR / "postamoldovei_clienti.txt"
OUTPUT_DIR = Path(__file__).parent.parent.parent / "invoices_logs"


def load_client_names() -> list:
    """Load client names from config file."""
    names = []
    try:
        if CLIENTS_FILE.exists():
            with open(CLIENTS_FILE, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith('#'):
                        names.append(line)
            logger.debug(f"Posta: loaded {len(names)} client names")
    except Exception as e:
        logger.error(f"Posta: failed to load client names: {e}")
    return names


def query_clients_via_ssh(ssh, client_names: list) -> list:
    """Query billing database for client data."""
    results = []
    
    db_user = BILLING_CONFIG['db_user']
    db_pass = BILLING_CONFIG['db_password']
    db_name = BILLING_CONFIG['db_name']
    
    for name in client_names:
        # Split name into parts (supports "Firstname Lastname" or "Lastname Firstname")
        parts = name.strip().split()
        
        if len(parts) >= 2:
            # Two or more words - search both combinations
            part1 = parts[0].replace("'", "''")
            part2 = ' '.join(parts[1:]).replace("'", "''")
            where_clause = f"""(
                (c.firstname LIKE '%{part1}%' AND c.lastname LIKE '%{part2}%') OR
                (c.firstname LIKE '%{part2}%' AND c.lastname LIKE '%{part1}%')
            )"""
        else:
            # Single word - search in both fields
            safe_name = name.replace("'", "''")
            where_clause = f"(c.firstname LIKE '%{safe_name}%' OR c.lastname LIKE '%{safe_name}%')"
        
        query = f"""
        SELECT 
            cfv.value AS Contul,
            CONCAT(c.firstname, ' ', c.lastname) AS NPP,
            c.address1 AS Adresa,
            COALESCE(SUM(i.total), 0) AS Suma
        FROM tblclients c
        JOIN tblcustomfieldsvalues cfv ON cfv.relid = c.id AND cfv.fieldid = 15
        LEFT JOIN tblinvoices i ON i.userid = c.id AND i.status = 'Unpaid'
        WHERE {where_clause}
        GROUP BY c.id, cfv.value, c.firstname, c.lastname, c.address1
        HAVING Suma > 0;
        """
        
        cmd = f'mysql -u {db_user} -p"{db_pass}" {db_name} -N -e "{query}"'
        
        stdin, stdout, stderr = ssh.exec_command(cmd)
        output = stdout.read().decode('utf-8', errors='ignore').strip()
        
        if output:
            for line in output.split('\n'):
                parts = line.split('\t')
                if len(parts) >= 4:
                    results.append({
                        'Contul': parts[0],
                        'NPP': parts[1],
                        'Adresa': parts[2],
                        'Suma': int(float(parts[3])) if parts[3] else 0,
                    })
    
    return results


def create_xlsx(data: list, output_path: Path) -> bool:
    """Create XLSX file with exact format for Posta Moldovei."""
    try:
        import openpyxl
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        # Headers
        ws['A1'] = 'Contul'
        ws['B1'] = 'NPP'
        ws['C1'] = 'Adresa'
        ws['D1'] = 'Suma'
        
        # Data
        for i, row in enumerate(data, start=2):
            ws[f'A{i}'] = row['Contul']
            ws[f'B{i}'] = row['NPP']
            ws[f'C{i}'] = row['Adresa']
            ws[f'D{i}'] = row['Suma']
        
        wb.save(output_path)
        logger.info(f"Posta: XLSX created: {output_path}")
        return True
        
    except Exception as e:
        logger.error(f"Posta: XLSX creation error: {e}")
        return False


def upload_to_ftp(file_path: Path) -> bool:
    """Upload file to Posta Moldovei FTP."""
    
    if not FTP_CONFIG["user"] or not FTP_CONFIG["password"]:
        logger.warning("Posta: FTP credentials not configured")
        return False
    
    try:
        from ftplib import FTP
        
        logger.info(f"Posta: connecting to FTP {FTP_CONFIG['host']}...")
        ftp = FTP(FTP_CONFIG['host'], timeout=60)
        ftp.login(FTP_CONFIG['user'], FTP_CONFIG['password'])
        
        if FTP_CONFIG['path']:
            ftp.cwd(FTP_CONFIG['path'])
        
        with open(file_path, 'rb') as f:
            ftp.storbinary(f'STOR {file_path.name}', f)
        
        ftp.quit()
        logger.info(f"Posta: file uploaded to FTP: {FTP_CONFIG['path']}{file_path.name}")
        return True
        
    except Exception as e:
        logger.error(f"Posta: FTP upload error: {e}")
        return False


def generate_posta_invoice() -> dict:
    """Main function: query billing, create XLSX, upload to FTP."""
    result = {
        "success": False,
        "message": "",
        "file": None,
        "records": 0,
        "timestamp": datetime.now().isoformat(),
    }
    
    ssh = None
    
    try:
        # 1. Load client names
        client_names = load_client_names()
        if not client_names:
            result["message"] = "No clients in postamoldovei_clienti.txt"
            return result
        
        logger.info(f"Posta: searching for {len(client_names)} clients...")
        
        # 2. Connect via SSH
        logger.info(f"Posta: connecting to {BILLING_CONFIG['ssh_host']}...")
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        ssh.connect(
            hostname=BILLING_CONFIG['ssh_host'],
            port=BILLING_CONFIG['ssh_port'],
            username=BILLING_CONFIG['ssh_user'],
            password=BILLING_CONFIG['ssh_password'],
            timeout=30,
        )
        
        # 3. Query database
        data = query_clients_via_ssh(ssh, client_names)
        
        if not data:
            result["message"] = "No clients with unpaid invoices found"
            result["success"] = True
            return result
        
        logger.info(f"Posta: found {len(data)} records")
        result["records"] = len(data)
        
        # 4. Create XLSX
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        filename = f"rapidlink_{datetime.now().strftime('%Y%m%d')}.xlsx"
        output_path = OUTPUT_DIR / filename
        
        if not create_xlsx(data, output_path):
            result["message"] = "Failed to create XLSX"
            return result
        
        result["file"] = str(output_path)
        
        # 5. Upload to FTP
        ftp_ok = upload_to_ftp(output_path)
        
        result["success"] = True
        if ftp_ok:
            result["message"] = f"Posta: {len(data)} records | FTP ✓"
        else:
            result["message"] = f"Posta: {len(data)} records | FTP ✗"
        
        return result
        
    except Exception as e:
        logger.error(f"Posta: error: {e}")
        result["message"] = f"Error: {str(e)}"
        return result
    finally:
        if ssh:
            ssh.close()
